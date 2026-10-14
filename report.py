"""Excel-отчёт: лист Follow-up в виде блоков (дата → пользователь → попытки) и листы сводок."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import rules

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
DATE_FILL = PatternFill("solid", fgColor="BFBFBF")
USER_FILL = PatternFill("solid", fgColor="BDD7EE")
ACTION_FILLS = {
    rules.RESOLVED: PatternFill("solid", fgColor="C6EFCE"),
    rules.CONTACT: PatternFill("solid", fgColor="FFC7CE"),
    rules.CRM: PatternFill("solid", fgColor="FFEB9C"),
    rules.ESCALATE: PatternFill("solid", fgColor="E4DFEC"),
}

FOLLOWUP_COLUMNS = [  # (заголовок, ширина)
    ("Transaction", 13), ("User", 10), ("Created", 21), ("State", 10), ("Status code", 32),
    ("Method", 13), ("Provider", 11), ("Amount", 9), ("Email", 26), ("Name", 9), ("Country", 8),
    ("Action", 21), ("Message to send", 60), ("Note", 20),
]


def _header(ws, columns):
    ws.append([title for title, _ in columns])
    for i, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def _fill_row(ws, fill, bold=False):
    for cell in ws[ws.max_row]:
        cell.fill = fill
        if bold:
            cell.font = Font(bold=True)


def _followup_sheet(wb, blocks):
    ws = wb.create_sheet("Follow-up")
    _header(ws, FOLLOWUP_COLUMNS)
    current_date = None
    for block in blocks:
        if block["date"] != current_date:
            current_date = block["date"]
            ws.append([f"── {current_date:%d.%m.%Y} ──"])
            _fill_row(ws, DATE_FILL, bold=True)

        first = block["attempts"][0]
        ws.append(["", block["user_id"], "", "", "", "", "", "", first["email"], first["first_name"], first["country"]])
        _fill_row(ws, USER_FILL, bold=True)

        for i, a in enumerate(block["attempts"]):
            is_last = i == len(block["attempts"]) - 1
            ws.append([a["id"], a["user_id"], a["created"].strftime("%Y-%m-%d %H:%M:%S"), a["state"],
                       a["status_code"], a["method"], a["provider"], a["amount"], "", "", "",
                       block["action"] if is_last else "",
                       block["message"] if is_last else "",
                       block["note"] if is_last else ""])
            if is_last:
                ws.cell(row=ws.max_row, column=12).fill = ACTION_FILLS[block["action"]]
        ws.append([])


def _table_sheet(wb, title, df, widths, percent_cols=()):
    ws = wb.create_sheet(title)
    _header(ws, list(zip(df.columns, widths)))
    for row in df.itertuples(index=False):
        ws.append(list(row))
    if "Date" in df.columns:
        for (cell,) in ws.iter_rows(min_row=2, max_col=1):
            cell.number_format = "DD.MM.YYYY"
    for col in percent_cols:
        idx = list(df.columns).index(col) + 1
        for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            cell.number_format = '0.0"%"'
    ws.auto_filter.ref = ws.dimensions


def save_report(blocks, summaries, path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    _followup_sheet(wb, blocks)

    daily = summaries["daily"].rename(columns={"date": "Date", "transactions": "Transactions", "users": "Users",
                                               "success_rate": "Success rate", "volume_eur": "Volume, EUR"})
    _table_sheet(wb, "Daily", daily, [12, 13, 8, 13, 13] + [21] * (len(daily.columns) - 5), ["Success rate"])

    providers = summaries["providers"].rename(columns={"method": "Method", "provider": "Provider",
                                                       "transactions": "Transactions", "success_rate": "Success rate"})
    _table_sheet(wb, "Providers", providers, [15, 13, 13, 13], ["Success rate"])

    errors = summaries["errors"].rename(columns={"status_code": "Status code", "count": "Count",
                                                 "users": "Users", "share": "Share"})
    _table_sheet(wb, "Errors", errors, [30, 9, 9, 9], ["Share"])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
